"""pptx_patch.py — Surgical patch of the existing PPTX.

Replaces the main image on Slides 3, 4, and 5 with upgraded figures produced
by the Tier 2 notebook implementation, then appends two new appendix slides
(D and E) with additional evidence plots.

Usage:
    uv run python src/her2rwd/pptx_patch.py

What is touched:
  - Slide 3 (index 2): largest picture replaced with q1_stacked_bars_by_stratum.png
  - Slide 4 (index 3): largest picture replaced with enhertu_uptake_ci.png
  - Slide 5 (index 4): largest picture replaced with sensitivity_dot_plot.png
  - New Appendix D appended: her2_timing.png (HER2 window justification)
  - New Appendix E appended: swimmer_plot.png + discordant_timelines.png

What is NOT touched:
  - All text boxes, bullet points, logos, and other shapes on existing slides
  - config/analysis_rules.yaml (rebuild_pptx stays false)
  - report.py
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches, Pt

# ── Paths ─────────────────────────────────────────────────────────────────
_ROOT = Path(__file__).resolve().parent.parent.parent
_PPTX = _ROOT / "reports" / "HER2_mBC_Enhertu_Analysis.pptx"
_FIGS = _ROOT / "outputs" / "figures"

# ── Slide dimensions ──────────────────────────────────────────────────────
_SLIDE_W = Inches(13.33)
_SLIDE_H = Inches(7.5)

# ── Colours ───────────────────────────────────────────────────────────────
_BLUE = RGBColor(0x00, 0x4B, 0x87)
_GREY = RGBColor(0x58, 0x58, 0x58)
_SILVER = RGBColor(0x99, 0x99, 0x99)

# ── Image substitutions: (0-based slide index, new figure filename) ───────
_SUBSTITUTIONS = [
    (2, "q1_stacked_bars_by_stratum.png"),  # Slide 3 Q1 — adds curated/EHR row
    (3, "enhertu_uptake_ci.png"),            # Slide 4 Q2 — adds Wilson CI error bars
    (4, "sensitivity_dot_plot.png"),         # Slide 5   — visual dot plot > table image
]


def _find_largest_picture(slide):
    """Return the largest PICTURE shape by pixel area, or None."""
    best, best_area = None, 0
    for shape in slide.shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            area = shape.width * shape.height
            if area > best_area:
                best_area = area
                best = shape
    return best


def _replace_largest_image(slide, new_path: Path) -> bool:
    """Swap the largest picture on the slide for new_path at the same bbox."""
    old = _find_largest_picture(slide)
    if old is None:
        print("    WARNING: no picture shape found — skipping")
        return False
    left, top, width, height = old.left, old.top, old.width, old.height
    old._element.getparent().remove(old._element)
    slide.shapes.add_picture(str(new_path), left, top, width, height)
    return True


def _add_textbox(slide, text: str, left, top, width, height,
                 size: int = 12, bold: bool = False,
                 color: RGBColor = None) -> None:
    color = color or _GREY
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    run = p.runs[0]
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def _append_appendix_slide(
    prs: Presentation,
    title: str,
    bullets: list[str],
    images: list[tuple[Path, float, float, float, float]],
    footer: str,
) -> None:
    """Add a blank slide with title, optional bullets, image(s), and footer."""
    # Blank layout — index 6 in most themes, fall back to last layout
    try:
        layout = prs.slide_layouts[6]
    except IndexError:
        layout = prs.slide_layouts[-1]

    slide = prs.slides.add_slide(layout)

    # Title
    _add_textbox(
        slide, title,
        left=Inches(0.3), top=Inches(0.1),
        width=Inches(12.7), height=Inches(0.65),
        size=15, bold=True, color=_BLUE,
    )

    # Bullets
    if bullets:
        bullet_text = "\n".join(f"•  {b}" for b in bullets)
        _add_textbox(
            slide, bullet_text,
            left=Inches(0.3), top=Inches(0.78),
            width=Inches(12.7), height=Inches(0.7),
            size=9, bold=False, color=_GREY,
        )

    # Images
    for img_path, l, t, w, h in images:
        if img_path.exists():
            slide.shapes.add_picture(str(img_path), Inches(l), Inches(t), Inches(w), Inches(h))
        else:
            print(f"    WARNING: {img_path.name} not found — image slot left empty")

    # Footer
    _add_textbox(
        slide, footer,
        left=Inches(0.3), top=Inches(7.08),
        width=Inches(12.7), height=Inches(0.35),
        size=7, bold=False, color=_SILVER,
    )


def main() -> None:
    if not _PPTX.exists():
        raise FileNotFoundError(f"PPTX not found: {_PPTX}")

    prs = Presentation(str(_PPTX))
    n_before = len(prs.slides)
    print(f"Loaded  {_PPTX.name}  ({n_before} slides)")

    # ── Replace images on Slides 3, 4, 5 ─────────────────────────────────
    for idx, fig_name in _SUBSTITUTIONS:
        fig = _FIGS / fig_name
        if not fig.exists():
            print(f"  Slide {idx + 1}: SKIP — {fig_name} not found")
            continue
        ok = _replace_largest_image(prs.slides[idx], fig)
        print(f"  Slide {idx + 1}: {'OK image replaced' if ok else 'XX no picture to replace'}  ->  {fig_name}")

    # ── Appendix D — HER2 test timing evidence ────────────────────────────
    _append_appendix_slide(
        prs,
        title=(
            "HER2 tests fall mostly AFTER the metastatic index date — "
            "the +90-day post-index window captures the clinical workup (ADR-007)"
        ),
        bullets=[
            "Most biomarker workup occurs post-staging, not pre-staging — window must extend after the index date",
            "Sensitivity check S1 (narrow −180/+90 d window) produces identical classification, confirming robustness",
        ],
        images=[
            (_FIGS / "her2_timing.png", 0.3, 1.55, 12.7, 5.2),
        ],
        footer=(
            "Fully synthetic data · Source: outputs/figures/her2_timing.png · "
            "Rule set v1 (ADR-007)"
        ),
    )
    print("  Appended Appendix D  (HER2 test timing justification)")

    # ── Appendix E — Patient timelines + discordant cases ─────────────────
    _append_appendix_slide(
        prs,
        title=(
            "Patient treatment journeys confirm 2L Enhertu dominance; "
            "discordant patients received it with no HER2 test in the biomarker window"
        ),
        bullets=[
            "Left: swimmer plot, n=15 sample patients. Orange bar = Enhertu; diamond = HER2 test (colour = HER2 category).",
            "Right: 3 illustrative discordant patients — Enhertu dispensed with no HER2 result in the −365/+90 d window (ADR-013).",
        ],
        images=[
            (_FIGS / "swimmer_plot.png",        0.3,  1.55, 7.8, 5.2),
            (_FIGS / "discordant_timelines.png", 8.3,  1.55, 4.75, 5.2),
        ],
        footer=(
            "Fully synthetic data · Source: outputs/figures/ · "
            "Rule set v1 (ADR-013)"
        ),
    )
    print("  Appended Appendix E  (swimmer + discordant timelines)")

    # ── Save ──────────────────────────────────────────────────────────────
    prs.save(str(_PPTX))
    n_after = len(prs.slides)
    print(f"\nSaved  {_PPTX.name}  ({n_before} -> {n_after} slides)")
    print("Verify: open PPTX and check slides 3, 4, 5 and new Appendix D/E.")


if __name__ == "__main__":
    main()
